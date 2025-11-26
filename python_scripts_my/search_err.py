import os
import sys
import openpyxl

def check_missing_files(root_folder):
    missing_folders = []
    
    for foldername, _, filenames in os.walk(root_folder):
        # 检查当前文件夹中是否有missHistMap.txt
        if 'missHistMap.txt' not in filenames:
            # 获取相对于根文件夹的路径
            rel_path = os.path.relpath(foldername, root_folder)
            missing_folders.append(rel_path)
    
    # 对结果进行排序（按字母顺序）
    missing_folders.sort()
    return missing_folders

def save_to_excel(missing_folders, output_file):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Missing Files"
    
    # 写入标题
    ws['A1'] = "序号"
    ws['B1'] = "Folder Path"
    
    # 写入数据（带序号）
    for idx, folder in enumerate(missing_folders, start=2):
        ws[f'A{idx}'] = idx - 1  # 序号
        ws[f'B{idx}'] = folder
    
    # 自动调整列宽
    for column in ['A', 'B']:
        max_length = 0
        for cell in ws[column]:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2) * 1.2
        ws.column_dimensions[column].width = adjusted_width
    
    wb.save(output_file)
    print(f"结果已保存到 {output_file}")

if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("用法: python script.py <文件夹路径>")
        sys.exit(1)
    
    folder_path = sys.argv[1]
    if not os.path.isdir(folder_path):
        print(f"错误: {folder_path} 不是一个有效的文件夹路径")
        sys.exit(1)
    
    missing = check_missing_files(folder_path)
    if missing:
        output_filename = "missing_folders.xlsx"
        save_to_excel(missing, output_filename)
        print(f"共找到 {len(missing)} 个缺失文件的文件夹")
    else:
        print("所有子文件夹中都包含 missHistMap.txt 文件")